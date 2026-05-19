import openpyxl
import os

upload_dir = os.path.join(os.path.dirname(__file__), 'data', 'uploads')
files = os.listdir(upload_dir)
for f in files:
    if f.endswith('.xlsx'):
        path = os.path.join(upload_dir, f)
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.active
        if ws is None:
            print(f"{f}: no active worksheet")
            continue
        print(f'File: {f}')
        print('Header row:')
        for cell in ws[1]:
            if cell.value:
                print(f'  col {cell.column}: {repr(cell.value)}')
        print(f'Total data rows: {ws.max_row - 1}')
        print()
